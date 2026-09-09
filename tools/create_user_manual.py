from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('docs/Victo OMS User Manual.docx')
NAVY = '172033'
BLUE = '3457D5'
GOLD = 'E8B55D'
PALE = 'F5F7FB'
LINE = 'D9E0EA'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr
    el = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = OxmlElement(f'w:{side}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '4')
        tag.set(qn('w:color'), LINE)
        el.append(tag)
    tblPr.append(el)

def cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p

def step(doc, number, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f'{number}. {title}. ')
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(body)

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    borders(table)
    for i, val in enumerate(rows[0]):
        shade(table.rows[0].cells[i], NAVY)
        cell_text(table.rows[0].cells[i], val, True, 'FFFFFF')
    for row_idx, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if row_idx % 2 == 1:
                shade(cells[i], PALE)
            cell_text(cells[i], val)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.7)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.72)
section.right_margin = Inches(.72)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
for name in ('Heading 1', 'Heading 2', 'Heading 3'):
    styles[name].font.name = 'Aptos Display'
    styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos Display')
styles['Heading 1'].font.size = Pt(16)
styles['Heading 1'].font.bold = True
styles['Heading 2'].font.size = Pt(12)
styles['Heading 2'].font.bold = True

# Header and footer
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run('VICTO OMS  |  USER MANUAL')
r.font.size = Pt(8)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string('68758B')
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Victo OMS  •  Internal operations guide  •  September 2026')
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string('68758B')

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(78)
r = p.add_run('V')
r.bold = True
r.font.size = Pt(38)
r.font.color.rgb = RGBColor.from_string(GOLD)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Victo OMS User Manual')
r.font.color.rgb = RGBColor(0, 0, 0)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Order management and production workflow guide')
r.font.size = Pt(14)
r.font.color.rgb = RGBColor.from_string('56647A')
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(28)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('For Owners, Administrators, Designers and Cameramen')
r.font.size = Pt(10)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(BLUE)
doc.add_page_break()

heading(doc, 'How to use this manual')
doc.add_paragraph('This manual explains the working process in Victo OMS from customer setup to final delivery or pickup. Use the section for your role first, then refer to the shared workflow when an order moves to the next team member.')
heading(doc, 'Contents', 2)
for item in ['1. System overview and roles', '2. Sign in and navigation', '3. Owner and administrator operations', '4. Designer workflow', '5. Cameraman workflow', '6. Order status guide', '7. Notifications and common checks']:
    bullet(doc, item)

heading(doc, '1. System overview and roles')
doc.add_paragraph('Victo OMS is the central workspace for managing a custom order after it has been received. It keeps customer details, order items, production progress, design files, photos, approvals and delivery activity in one place.')
add_table(doc, [
    ['Role', 'Main responsibility', 'Key workspace'],
    ['Owner', 'Oversees operations, approves designs, manages delivery and reports.', 'Dashboard, Orders, Customers, Monitoring, Reports'],
    ['Administrator', 'Creates and maintains orders and customers; supports operational monitoring.', 'Admin Workspace, Order Management, Customer Management'],
    ['Designer', 'Starts assigned work, uploads design files and submits for approval.', 'My Tasks'],
    ['Cameraman', 'Runs photo sessions and uploads product photos.', 'Photo Tasks'],
])

heading(doc, '2. Sign in and navigation')
step(doc, 1, 'Open Victo OMS', 'Go to the Victo OMS web address provided by your organisation.')
step(doc, 2, 'Sign in', 'Enter your registered email address and password, then select Log in.')
step(doc, 3, 'Use the sidebar', 'Your role determines the menu items you can access. Select a menu item to open the related workspace.')
step(doc, 4, 'Check notifications', 'Select the bell in the top header to see recent assignments, approvals, revisions and photo-session updates. Select a notification to mark it as read.')
step(doc, 5, 'Sign out safely', 'Use the Logout button in the top-right corner when you finish working, especially on shared computers.')
doc.add_page_break()

heading(doc, '3. Owner and administrator operations')
heading(doc, 'Create a customer', 2)
step(doc, 1, 'Open Customers', 'Select Customers from the Owner menu, or Customer Management from the Administrator menu.')
step(doc, 2, 'Add the record', 'Select Add Customer or Create Customer, complete the contact details and save. Use a clear customer name and a valid phone number so the order team can contact the customer.')
step(doc, 3, 'Maintain records', 'Use View to review a customer, Edit to correct details, and Delete only when the record is no longer needed and it is safe to remove.')

heading(doc, 'Create and manage an order', 2)
step(doc, 1, 'Open the order form', 'Choose Create Order from the sidebar or Orders page.')
step(doc, 2, 'Select or add a customer', 'Choose the customer record. Use the quick-customer option if the person has not yet been added.')
step(doc, 3, 'Enter order details', 'Add the order items, quantities, prices, due date, delivery method and customer brief. Attach reference files when provided by the customer.')
step(doc, 4, 'Save the order', 'Review the information before saving. The system creates the order record and it becomes available for the next workflow action.')
step(doc, 5, 'Find an existing order', 'Use Search, Sort By and Direction on the Orders page. Open View Order to see the full timeline, files, items and available actions.')

heading(doc, 'Review a submitted design', 2)
doc.add_paragraph('Only the Owner can approve a design or request a revision.')
step(doc, 1, 'Open the order', 'Find an order with the Pending Approval status and select View Order.')
step(doc, 2, 'Review files and brief', 'Check the uploaded design against the order items, customer brief and reference files.')
step(doc, 3, 'Choose an outcome', 'Select Approve when the design is ready for production. Select Request Revision when changes are needed, and provide clear feedback for the designer.')

heading(doc, 'Production, photo and fulfilment actions', 2)
add_table(doc, [
    ['Action', 'When to use it', 'Result'],
    ['Ready at HQ', 'The order is ready for the HQ or photo stage.', 'Order becomes available for cameraman assignment.'],
    ['Assign Cameraman', 'A photo session is required.', 'The selected cameraman receives the task.'],
    ['Dispatch Delivery', 'An order is being sent to the customer.', 'Status changes to Out for Delivery.'],
    ['Ready for Pickup', 'The customer should collect the order.', 'Status changes to Waiting for Pickup.'],
    ['Mark Delivered or Confirm Pickup', 'The order has reached the customer.', 'Order is completed.'],
])
doc.add_page_break()

heading(doc, '4. Designer workflow')
heading(doc, 'Work on an assigned task', 2)
step(doc, 1, 'Open My Tasks', 'Select My Tasks from the sidebar. The dashboard also shows current tasks, pending approvals and overdue work.')
step(doc, 2, 'Review the order', 'Open the assigned order and read the customer brief, product details, due date and reference files before starting.')
step(doc, 3, 'Start the task', 'Select Start Task. This changes the working state to In Progress so the team can see that design work has begun.')
step(doc, 4, 'Upload design files', 'Use the upload section to attach the correct design files. Confirm that the file name makes the version easy to identify.')
step(doc, 5, 'Submit for approval', 'When the design is ready, select Submit for Approval. The owner is notified and the order moves to Pending Approval.')
step(doc, 6, 'Respond to revisions', 'If a revision is requested, read the feedback on the order, update the files and submit again for approval.')

heading(doc, 'Create a job order', 2)
doc.add_paragraph('When the approved work needs a production job order, open the relevant order and choose the job-order option. Enter the production instructions accurately, save the job order, then use Generate Word when a printable copy is needed.')

heading(doc, 'Designer good practice', 2)
bullet(doc, 'Start work only after checking the customer brief and all reference assets.')
bullet(doc, 'Use clear file names that include the order number and version where appropriate.')
bullet(doc, 'Do not submit incomplete designs. Confirm the item count, sizing and copy before submitting.')
bullet(doc, 'Watch due dates and the notifications bell throughout the day.')

heading(doc, '5. Cameraman workflow')
step(doc, 1, 'Open Photo Tasks', 'Select Photo Tasks from the sidebar to see orders waiting at HQ or already in a photo session.')
step(doc, 2, 'Open the task', 'Review the order details and any customer or product instructions before the session begins.')
step(doc, 3, 'Start Photo Session', 'Select Start Photo Session when shooting begins. The order status changes to Photo Session.')
step(doc, 4, 'Upload product photos', 'Upload the final usable photos to the order. Check that each upload belongs to the correct order.')
step(doc, 5, 'Complete Photo Session', 'Select Complete only after all required images have been uploaded. The order becomes Photo Completed and returns to the operations team for fulfilment.')
doc.add_page_break()

heading(doc, '6. Order status guide')
add_table(doc, [
    ['Status', 'Meaning', 'Typical next action'],
    ['Pending', 'Order has been created and awaits assignment or work.', 'Assign designer or prepare the order.'],
    ['Assigned', 'A designer has been assigned.', 'Designer starts the task.'],
    ['In Progress', 'Design work is underway.', 'Designer uploads and submits design.'],
    ['Pending Approval', 'Design is ready for owner review.', 'Owner approves or requests revision.'],
    ['Printing', 'Approved work is in production.', 'Prepare for HQ or photo stage.'],
    ['Ready at HQ', 'Order is ready at headquarters.', 'Assign cameraman if a session is needed.'],
    ['Photo Session', 'Photography is underway.', 'Cameraman uploads photos and completes session.'],
    ['Photo Completed', 'Photography is complete.', 'Arrange delivery or pickup.'],
    ['Out for Delivery', 'Order is with delivery.', 'Mark delivered after confirmation.'],
    ['Waiting for Pickup', 'Order is ready for collection.', 'Confirm pickup after collection.'],
    ['Completed', 'Order has been fulfilled.', 'No further workflow action required.'],
])

heading(doc, '7. Notifications and common checks')
heading(doc, 'Before moving an order forward', 2)
bullet(doc, 'Confirm you are viewing the correct order number and customer.')
bullet(doc, 'Check the due date and current status.')
bullet(doc, 'Read the customer brief and review attachments before approving, designing or shooting.')
bullet(doc, 'Make sure uploaded files and photos are complete and belong to the correct order.')
bullet(doc, 'Use the status action only when the related work has genuinely been completed.')

heading(doc, 'If something is missing or incorrect', 2)
doc.add_paragraph('Do not move the order to the next stage until the information is clear. Review the order timeline, files and notes first. For design changes, use the revision workflow so the designer receives a visible instruction. For an incorrect customer or order record, use Edit before downstream work begins.')

heading(doc, 'Need help', 2)
doc.add_paragraph('If you cannot see a menu item or action, confirm that you are signed in with the correct role. If the issue continues, record the order number, the action you attempted and any message shown on screen, then contact the system administrator.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
